import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from multi_agent_docs.config import Settings
from multi_agent_docs.models import LLMAnswer, TokenCount


@pytest.fixture
def docx_payload() -> bytes:
    document = Document()
    document.add_paragraph("Начало документа")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Товар"
    table.cell(0, 1).text = "Цена"
    table.cell(1, 0).text = "Книга"
    table.cell(1, 1).text = "350"
    document.add_paragraph("Конец документа")
    target = BytesIO()
    document.save(target)
    return target.getvalue()


class FakeLLM:
    def __init__(self, invalid_fact_checks: int = 0):
        self.invalid_fact_checks = invalid_fact_checks

    def count_tokens(self, text: str) -> int:
        return max(1, len(text.split()))

    def ask(self, prompt: str, *, system_prompt: str = "") -> LLMAnswer:
        usage = TokenCount(input_tokens=10, output_tokens=5, total_tokens=15)
        if "Верни только JSON" in prompt or "Исправь ответ" in prompt:
            if self.invalid_fact_checks:
                self.invalid_fact_checks -= 1
                return LLMAnswer(content="not json", usage=usage)
            return LLMAnswer(
                content=json.dumps(
                    {
                        "verification_summary": "Проверено",
                        "doubtful_claims": [],
                        "contradictions": [],
                        "neuroslop": {
                            "probability": 0.25,
                            "explanation": "Низкая доля шаблонных формулировок",
                        },
                    },
                    ensure_ascii=False,
                ),
                usage=usage,
            )
        return LLMAnswer(content="Результат анализа", usage=usage)


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(
        app_mode="mock",
        database_path=tmp_path / "app.db",
        upload_dir=tmp_path / "jobs",
        max_upload_bytes=100 * 1024 * 1024,
        max_document_tokens=250_000,
        chunk_target_tokens=8_000,
        chunk_overlap_tokens=300,
        gigachat_credentials="test",
        gigachat_scope="test",
        gigachat_model="test",
        gigachat_verify_ssl=True,
    )
