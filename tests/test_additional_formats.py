from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from multi_agent_docs.errors import CorruptDocumentError, DocumentTooLargeError, EmptyDocumentError
from multi_agent_docs.ingestion.loader import load_document


def test_docx_keeps_paragraphs_and_tables_in_document_order(tmp_path: Path, docx_payload):
    path = tmp_path / "report.DOCX"
    path.write_bytes(docx_payload)
    text = load_document(str(path))
    assert text.index("Начало документа") < text.index("Товар")
    assert "Товар\tЦена" in text
    assert "Книга\t350" in text
    assert text.index("350") < text.index("Конец документа")


def test_docx_reads_nested_tables(tmp_path: Path):
    document = Document()
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "Внешняя ячейка"
    nested = outer.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Вложенные данные"
    path = tmp_path / "nested.docx"
    document.save(path)
    text = load_document(str(path))
    assert text.index("Внешняя ячейка") < text.index("Вложенные данные")


def test_docx_rejects_empty_document(tmp_path: Path):
    path = tmp_path / "empty.docx"
    Document().save(path)
    with pytest.raises(EmptyDocumentError):
        load_document(str(path))


@pytest.mark.parametrize("payload", [b"not a DOCX", b"PK\x03\x04truncated archive"])
def test_docx_rejects_corrupt_files(tmp_path: Path, payload):
    path = tmp_path / "broken.docx"
    path.write_bytes(payload)
    with pytest.raises(CorruptDocumentError):
        load_document(str(path))


def test_docx_rejects_other_zip_archives(tmp_path: Path):
    path = tmp_path / "renamed.docx"
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as package:
        package.writestr("notes.txt", "Это не Word")
    with pytest.raises(CorruptDocumentError):
        load_document(str(path))


def test_docx_limits_unpacked_size(tmp_path: Path, docx_payload, monkeypatch):
    path = tmp_path / "large.docx"
    path.write_bytes(docx_payload)
    monkeypatch.setattr("multi_agent_docs.ingestion.additional_formats._MAX_DOCX_UNPACKED_BYTES", 100)
    with pytest.raises(DocumentTooLargeError):
        load_document(str(path))


def test_html_extracts_content_without_script_style_or_head(tmp_path: Path):
    path = tmp_path / "page.HTM"
    path.write_text(
        '<html><head><title>Служебный заголовок</title><style>hidden_css</style></head>'
        '<body><h1>Заголовок статьи</h1><p>Текст <strong>важной</strong> статьи &amp; заметки.</p>'
        '<script>secret_script()</script><noscript>Запасной текст</noscript>'
        '<table><tr><td>Книга</td><td>350</td></tr></table>'
        '<a href="https://example.com/source">Источник</a>'
        '<img src="image.png" alt="Подпись изображения"></body></html>',
        encoding="utf-8",
    )
    text = load_document(str(path))
    assert "Заголовок статьи" in text
    assert "Текст важной статьи & заметки." in text
    assert "Книга\t350" in text
    assert "Источник (https://example.com/source)" in text
    assert "Подпись изображения" in text
    assert "Служебный заголовок" not in text
    assert "hidden_css" not in text
    assert "secret_script" not in text
    assert "Запасной текст" not in text


def test_html_respects_declared_cyrillic_encoding(tmp_path: Path):
    path = tmp_path / "legacy.html"
    path.write_bytes('<meta charset="windows-1251"><p>Привет, мир!</p>'.encode("cp1251"))
    assert "Привет, мир!" in load_document(str(path))


def test_html_accepts_document_fragment(tmp_path: Path):
    path = tmp_path / "fragment.html"
    path.write_text("<p>Первый абзац</p><p>Второй абзац</p>", encoding="utf-8")
    text = load_document(str(path))
    assert "Первый абзац\n" in text
    assert text.endswith("Второй абзац")


def test_html_table_keeps_empty_columns_despite_formatting_whitespace(tmp_path: Path):
    path = tmp_path / "table.html"
    path.write_text(
        '<p>Начало</p><table><tr>\n<td></td>\n<td>Книга</td>\n<td>350</td>\n</tr></table><p>Конец</p>',
        encoding="utf-8",
    )
    assert "\n\n\tКнига\t350\n\n" in load_document(str(path))


def test_markdown_keeps_meaning_tables_links_and_code(tmp_path: Path):
    path = tmp_path / "notes.MARKDOWN"
    path.write_text(
        '# Заголовок\n\nЭто **важный** текст и [источник](https://example.com).\n\n'
        '- Первый пункт\n- Второй пункт\n\n'
        '| Товар | Цена |\n| --- | --- |\n| Книга | 350 |\n\n'
        '```python\ndef example():\n    return "<tag>"\n```\n'
        '<script>secret_script()</script>',
        encoding="utf-8",
    )
    text = load_document(str(path))
    assert "Заголовок" in text
    assert "Это важный текст" in text
    assert "источник (https://example.com)" in text
    assert "Первый пункт" in text and "Второй пункт" in text
    assert "Книга\t350" in text
    assert '    return "<tag>"' in text
    assert "secret_script" not in text
    assert "**" not in text and "```" not in text


def test_markdown_preserves_code_indentation_at_document_edges(tmp_path: Path):
    path = tmp_path / "code.md"
    path.write_text('```python\n    return "Привет"\n```\n', encoding="utf-8")
    assert load_document(str(path)) == '    return "Привет"'


@pytest.mark.parametrize("delimiter", [",", ";", "\t", "|"])
def test_csv_detects_delimiter_and_keeps_cells(tmp_path: Path, delimiter):
    path = tmp_path / "export.CSV"
    path.write_text(f"Товар{delimiter}Цена\r\nКнига{delimiter}350\r\n", encoding="utf-8-sig")
    assert load_document(str(path)) == "Товар\tЦена\nКнига\t350"


def test_csv_handles_quoted_delimiters_newlines_and_empty_cells(tmp_path: Path):
    path = tmp_path / "quoted.csv"
    path.write_text('Имя,Описание,Цена\r\nКнига,"Первая строка, с запятой\nВторая строка",\r\n', encoding="utf-8")
    text = load_document(str(path))
    assert "Первая строка, с запятой\nВторая строка" in text
    assert text.startswith("Имя\tОписание\tЦена\nКнига\t")
    assert text.endswith("Вторая строка\t")


def test_csv_preserves_leading_empty_columns_and_quoted_whitespace(tmp_path: Path):
    path = tmp_path / "empty-column.csv"
    path.write_text('Имя,Описание\n,"  Пробелы внутри кавычек  "\n', encoding="utf-8")
    assert load_document(str(path)) == "Имя\tОписание\n\t  Пробелы внутри кавычек  "


def test_csv_accepts_single_column(tmp_path: Path):
    path = tmp_path / "names.csv"
    path.write_text("Имя\nЕгор\nАнтон\n", encoding="utf-8")
    assert load_document(str(path)) == "Имя\nЕгор\nАнтон"


def test_csv_rejects_unclosed_quotes(tmp_path: Path):
    path = tmp_path / "invalid.csv"
    path.write_text('Товар,Цена\n"Книга,350\n', encoding="utf-8")
    with pytest.raises(CorruptDocumentError):
        load_document(str(path))


@pytest.mark.parametrize(
    ("extension", "payload"),
    [("md", b"\xef\xbb\xbf\n"), ("md", b"---\n"),
     ("html", b"<html><head><style>a{}</style></head><body><script>x()</script></body></html>"),
     ("csv", b"\xef\xbb\xbf\n"), ("csv", b'"",""\n')],
)
def test_rejects_documents_without_text(tmp_path: Path, extension, payload):
    path = tmp_path / f"empty.{extension}"
    path.write_bytes(payload)
    with pytest.raises(EmptyDocumentError):
        load_document(str(path))


@pytest.mark.parametrize("extension", ["html", "md", "csv"])
def test_rejects_binary_payloads_in_text_formats(tmp_path: Path, extension):
    path = tmp_path / f"binary.{extension}"
    path.write_bytes(b"\x00binary")
    with pytest.raises(CorruptDocumentError):
        load_document(str(path))


@pytest.mark.parametrize("extension", ["md", "csv"])
def test_rejects_invalid_utf8(tmp_path: Path, extension):
    path = tmp_path / f"invalid.{extension}"
    path.write_bytes(b"\xff\xfe")
    with pytest.raises(CorruptDocumentError):
        load_document(str(path))
